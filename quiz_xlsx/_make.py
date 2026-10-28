# quiz_xlsx/_make.py — 產生 期中 複習 50 題 的 匯入檔（python3 quiz_xlsx/_make.py）
# 每一份：Kahoot（.xlsx）、Blooket（.csv）、Gimkit（.csv）、總表（.xlsx：題目＋答案＋Wordwall 複製用）＋大字 網頁 listen/2026-10-28_X/
import csv, os, re, sys, html
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _data import SETS

D = os.path.dirname(os.path.abspath(__file__))
TIME = 20
BAD = []
# 考卷、課本、網站 用過 的 人名（g3mid/_build.js、midterm/_build.js 的 BAN）
BAN = re.compile(r"\b(Mike|Emma|Wendy|Roy|Sam|Ann|Lily|Ken|Alan|Harry|Amy|Mia|Tom|Frank|Lisa|Peter|Grace|Owen|Paul|Tina)\b")
one = lambda q: q.replace("\n", "   ")

def pos_of(i):   # 正解 位置 平均 分到 1～4（不 照 順序，學生 猜不到）
    return [1, 3, 2, 4, 3, 1, 4, 2][i % 8]

ALL = {}
for name, L in SETS:
    if len(L) != 50: BAD.append(f"{name} 不是 50 題（{len(L)}）")
    seen = set(); rows = []
    for i, (q, ok, wr) in enumerate(L):
        opts = list(wr); p = pos_of(i); opts.insert(p - 1, ok)
        if len(opts) != 4 or len(set(opts)) != 4: BAD.append(f"{name} 第 {i+1} 題 選項 不是 4 個 不一樣 的")
        k = (q, ok)
        if k in seen: BAD.append(f"{name} 第 {i+1} 題 重複")
        seen.add(k)
        if len(one(q)) > 120: BAD.append(f"{name} 第 {i+1} 題 題目 太長")
        for o in opts:
            if len(o) > 75: BAD.append(f"{name} 第 {i+1} 題 選項 太長")
        t = q + " " + " ".join(opts)
        if BAN.search(t): BAD.append(f"{name} 第 {i+1} 題 用了 用過 的 人名 {BAN.search(t).group(0)}")
        if re.search(r"[一-鿿]", t): BAD.append(f"{name} 第 {i+1} 題 有 中文")
        rows.append((one(q), q, opts, p))
    ALL[name] = rows
# 同 年級 兩份 完全 不同（題目＋正解 一樣 ＝ 同一題）
for g in ("三年級", "四年級"):
    a, b = [set((r[0], r[2][r[3]-1]) for r in ALL[n]) for n in ALL if n.startswith(g)]
    if a & b: BAD.append(f"{g} 兩份 有 一樣 的 題目：{a & b}")
if BAD: print("✗ " + "\n✗ ".join(BAD)); sys.exit(1)

for name, rows in ALL.items():
    # Kahoot（官方 範本：第 8 列 標題、第 9 列 開始 題目）
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws["B2"] = "Quiz template"; ws["B3"] = name + "（50 題、每題 20 秒）"
    ws.append([]); ws.append([]); ws.append([]); ws.append([])
    ws.cell(8, 1, "#"); hd = ["Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters", "Answer 3 - max 75 characters", "Answer 4 - max 75 characters", "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs", "Correct answer(s) - choose at least one"]
    for c, h in enumerate(hd, 2): ws.cell(8, c, h)
    for i, (q, _, o, p) in enumerate(rows):
        r = 9 + i; ws.cell(r, 1, i + 1); ws.cell(r, 2, q)
        for k in range(4): ws.cell(r, 3 + k, o[k])
        ws.cell(r, 7, TIME); ws.cell(r, 8, str(p))
    wb.save(os.path.join(D, name + "_Kahoot.xlsx"))
    # Blooket（官方 CSV 範本）
    with open(os.path.join(D, name + "_Blooket.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Blooket\nImport Template", "", "", "", "", "", "", ""])
        w.writerow(["Question #", "Question Text", "Answer 1", "Answer 2", "Answer 3\n(Optional)", "Answer 4\n(Optional)", "Time Limit (sec)\n(Max: 300 seconds)", "Correct Answer(s)\n(Only include Answer #)"])
        for i, (q, _, o, p) in enumerate(rows): w.writerow([i + 1, q] + o + [TIME, p])
    # Gimkit（Spreadsheet Import：題目、正解、3 個 錯 的）
    with open(os.path.join(D, name + "_Gimkit.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Question", "Correct Answer", "Incorrect Answer 1", "Incorrect Answer 2 (Optional)", "Incorrect Answer 3 (Optional)"])
        for q, _, o, p in rows: w.writerow([q, o[p - 1]] + [x for k, x in enumerate(o) if k != p - 1])

# 總表（老師 看、Wordwall 複製 貼上）
wb = Workbook(); wb.remove(wb.active)
big = Font(name="Andika", size=16); bold = Font(name="Andika", size=16, bold=True, color="006100")
for name, rows in ALL.items():
    ws = wb.create_sheet(name.replace("_期中複習_", " "))
    ws.append(["#", "Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Correct", "Correct answer"])
    for i, (q, _, o, p) in enumerate(rows): ws.append([i + 1, q] + o + [p, o[p - 1]])
    for row in ws.iter_rows():
        for c in row: c.font = big; c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, 2 + ws.cell(r, 7).value); c.font = bold; c.fill = PatternFill("solid", fgColor="C6EFCE")
    for col, wd in zip("ABCDEFGH", (5, 46, 26, 26, 26, 26, 9, 26)): ws.column_dimensions[col].width = wd
    # Wordwall：一題 一行「題目 → 正解 → 錯 錯 錯」（Wordwall 的 Quiz 範本 可以 一格 一格 貼）
    ws2 = wb.create_sheet(name.replace("_期中複習_", " ") + " Wordwall")
    ws2.append(["Question", "Correct answer", "Wrong 1", "Wrong 2", "Wrong 3"])
    for q, _, o, p in rows: ws2.append([q, o[p - 1]] + [x for k, x in enumerate(o) if k != p - 1])
    for row in ws2.iter_rows():
        for c in row: c.font = big
    for col, wd in zip("ABCDE", (46, 26, 26, 26, 26)): ws2.column_dimensions[col].width = wd
wb.save(os.path.join(D, "期中複習_50題_全部_總表.xlsx"))

# 大字 網頁（老師 看 題目＋下載）
R = os.path.join(D, "..", "listen", "2026-10-28_X"); os.makedirs(R, exist_ok=True)
e = html.escape
sec = []
for name, rows in ALL.items():
    li = "".join(f'<li><div class="q">{e(q).replace(chr(10), "<br>")}</div><div class="o">' + "".join(f'<span class="{"ok" if k == p - 1 else ""}">{"①②③④"[k]} {e(x)}</span>' for k, x in enumerate(o)) + "</div></li>" for _, q, o, p in rows)
    dl = " ".join(f'<a href="../../quiz_xlsx/{e(name)}_{t}">⬇ {t.split(".")[0]}</a>' for t in ("Kahoot.xlsx", "Blooket.csv", "Gimkit.csv"))
    sec.append(f'<section><h2>{e(name.replace("_", " "))}</h2><p class="dl">{dl}</p><ol>{li}</ol></section>')
page = f'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>期中複習 50題</title><style>
@font-face{{font-family:Andika;src:url(../../words/fonts/andika-400.woff2) format("woff2");font-weight:400}}
@font-face{{font-family:Andika;src:url(../../words/fonts/andika-700.woff2) format("woff2");font-weight:700}}
:root{{--bg:#fffdf6;--fg:#1d1d1d;--ok:#0a7a2f;--okb:#d9f7e2;--card:#fff;--line:#ddd}}
@media (prefers-color-scheme:dark){{:root:not([data-theme="light"]){{--bg:#16181d;--fg:#f2f2f2;--ok:#7be39c;--okb:#173b24;--card:#20242b;--line:#3a3f48}}}}
:root[data-theme="dark"]{{--bg:#16181d;--fg:#f2f2f2;--ok:#7be39c;--okb:#173b24;--card:#20242b;--line:#3a3f48}}
body{{margin:0;padding:16px;background:var(--bg);color:var(--fg);font:22px/1.5 Andika,"Microsoft JhengHei",sans-serif}}
h1{{font-size:34px;margin:8px 0}} h2{{font-size:30px;margin:28px 0 8px}} .dl a{{display:inline-block;margin:6px 10px 6px 0;padding:10px 18px;border-radius:12px;background:#2f6fde;color:#fff;text-decoration:none;font-weight:700}}
.all a{{background:#c47a00}} ol{{padding-left:2.2em}} li{{margin:0 0 14px;padding:10px 12px;background:var(--card);border:1px solid var(--line);border-radius:12px}}
.q{{font-weight:700}} .o span{{display:inline-block;margin:4px 14px 0 0}} .o .ok{{color:var(--ok);background:var(--okb);font-weight:700;padding:0 8px;border-radius:8px}}
</style></head><body><h1>📝 期中複習 50 題（四選一）</h1>
<p>三年級 2 份、四年級 2 份；每一份 50 題、題目 完全 不同；綠色 ＝ 正確 答案。聽力 題 這次 不放。</p>
<p class="dl all"><a href="../../quiz_xlsx/期中複習_50題_全部_總表.xlsx">⬇ 全部 總表（含 Wordwall）</a></p>
{"".join(sec)}</body></html>'''
open(os.path.join(R, "index.html"), "w", encoding="utf-8").write(page)
from collections import Counter
for n, rows in ALL.items(): print(n, "50 題，正解 位置", dict(sorted(Counter(r[3] for r in rows).items())))
print("✓ 0 失敗")
